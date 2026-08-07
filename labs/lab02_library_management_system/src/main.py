"""Library Management System - Excel workbook processor and CLI entry point."""

from argparse import ArgumentParser
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

try:
    from .processor import determine_status, validate_book_record
except ImportError:
    from processor import determine_status, validate_book_record


LAB_DIR = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_FILE = LAB_DIR / "data" / "input" / "Library_Records_Input.xlsx"
DEFAULT_OUTPUT_FILE = LAB_DIR / "data" / "output" / "Library_Records_Updated.xlsx"


@dataclass
class ProcessingSummary:
    """Summary of the workbook processing results."""

    total_records: int = 0
    valid_records: int = 0
    invalid_records: int = 0
    errors: list[dict[str, object]] = field(default_factory=list)


def process_workbook(input_file=DEFAULT_INPUT_FILE, output_file=DEFAULT_OUTPUT_FILE):
    input_path = Path(input_file)
    output_path = Path(output_file)

    wb = load_workbook(input_path)
    ws = wb.active
    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary = ProcessingSummary(total_records=max(ws.max_row - 1, 0))

    for row in range(2, ws.max_row + 1):
        book_id = ws.cell(row, 1).value
        try:
            available = validate_book_record(
                ws.cell(row, 5).value,
                ws.cell(row, 6).value,
                ws.cell(row, 7).value,
            )
            ws.cell(row, 8).value = determine_status(available)
            summary.valid_records += 1
        except ValueError as err:
            message = str(err)
            ws.cell(row, 8).value = None
            summary.invalid_records += 1
            summary.errors.append(
                {"row": row, "book_id": book_id, "message": message}
            )

    wb.save(output_path)
    return summary


def build_parser():
    """Build the argument parser for the CLI."""
    parser = ArgumentParser(
        description="Process library book records from an Excel workbook."
    )
    parser.add_argument(
        "--input",
        default=DEFAULT_INPUT_FILE,
        type=Path,
        help="Path to the input Excel workbook.",
    )
    parser.add_argument(
        "--output",
        default=DEFAULT_OUTPUT_FILE,
        type=Path,
        help="Path where the updated workbook will be saved.",
    )
    return parser


def main():
    """Main entry point for the CLI."""
    args = build_parser().parse_args()
    summary = process_workbook(args.input, args.output)

    print("Library Management System")
    print(f"Input workbook : {args.input}")
    print(f"Output workbook: {args.output}")
    print(f"Total records  : {summary.total_records}")
    print(f"Status updated : {summary.valid_records}")
    print(f"Invalid records: {summary.invalid_records}")

    if summary.errors:
        print("Validation errors:")
        for error in summary.errors:
            print(f"  Row {error['row']} ({error['book_id']}): {error['message']}")

    return summary


if __name__ == "__main__":
    main()