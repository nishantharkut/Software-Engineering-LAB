"""Integration tests for the Excel processing pipeline."""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.main import (
    DEFAULT_INPUT_FILE,
    DEFAULT_OUTPUT_FILE,
    LAB_DIR,
    process_workbook,
)


def build_test_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "LibraryRecords"
    ws.append(
        [
            "Book ID",
            "Book Title",
            "Author",
            "Category",
            "Total Copies",
            "Available Copies",
            "Issued Copies",
            "Status",
        ]
    )
    # Valid record: Available
    ws.append(["B001", "Clean Code", "Robert C. Martin", "Programming", 10, 6, 4, None])
    # Valid record: Issued Out
    ws.append(["B002", "Design Patterns", "GoF", "Software Engineering", 5, 0, 5, None])
    # Invalid: constraint violation
    ws.append(["B003", "Invalid Book", "Author", "Category", 10, 7, 5, None])
    # Invalid: missing input
    ws.append(["B004", "Missing Field", "Author", "Category", 10, None, 5, None])
    # Invalid: text instead of number
    ws.append(["B005", "Text Field", "Author", "Category", 10, "Five", 5, None])
    # Invalid: negative inventory
    ws.append(["B006", "Negative Value", "Author", "Category", 10, -2, 12, None])
    wb.save(path)


def test_excel_pipeline(tmp_path):
    """Test the full Excel processing pipeline."""
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    build_test_workbook(input_file)

    summary = process_workbook(str(input_file), str(output_file))

    wb = load_workbook(output_file)
    ws = wb.active

    assert summary.total_records == 6
    assert summary.valid_records == 2
    assert summary.invalid_records == 4

    # Row 2: Valid - Available (6 copies)
    assert ws.cell(2, 8).value == "Available"

    # Row 3: Valid - Issued Out (0 copies)
    assert ws.cell(3, 8).value == "Issued Out"

    # Row 4: Invalid - constraint violation (Status remains blank)
    assert ws.cell(4, 8).value is None
    assert "does not equal" in summary.errors[0]["message"]

    # Row 5: Invalid - missing input (Status remains blank)
    assert ws.cell(5, 8).value is None
    assert summary.errors[1]["message"] == "Missing input"

    # Row 6: Invalid - text input (Status remains blank)
    assert ws.cell(6, 8).value is None
    assert summary.errors[2]["message"] == "Invalid data type"

    # Row 7: Invalid - negative inventory (Status remains blank)
    assert ws.cell(7, 8).value is None
    assert summary.errors[3]["message"] == "Negative inventory value"


def test_default_workbook_paths_are_anchored_to_lab_directory():
    """Verify default paths are relative to the lab directory."""
    assert (
        DEFAULT_INPUT_FILE
        == LAB_DIR / "data" / "input" / "Library_Records_Input.xlsx"
    )
    assert (
        DEFAULT_OUTPUT_FILE
        == LAB_DIR / "data" / "output" / "Library_Records_Updated.xlsx"
    )