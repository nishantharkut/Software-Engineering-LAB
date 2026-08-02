from openpyxl import Workbook, load_workbook

from labs.lab01_student_grade_processing.src.main import (
    DEFAULT_INPUT_FILE,
    DEFAULT_OUTPUT_FILE,
    LAB_DIR,
    process_workbook,
)


def build_input(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Records"
    ws.append(
        [
            "Roll Number",
            "Student Name",
            "Department",
            "Subject",
            "Project Marks (20)",
            "Minor Marks (20)",
            "Laboratory Marks (20)",
            "End Examination Marks (40)",
            "Grade",
        ]
    )
    ws.append(
        [
            "2026CS006",
            "Ananya Bose",
            "CSE",
            "Software Engineering",
            14,
            16,
            15,
            15,
            None,
        ]
    )
    ws.append(
        [
            "2026IT007",
            "Kabir Khan",
            "IT",
            "Software Engineering",
            20,
            20,
            20,
            41,
            None,
        ]
    )
    ws.append(
        [
            "2026ECE008",
            "Pooja Nair",
            "ECE",
            "Software Engineering",
            17,
            None,
            18,
            36,
            None,
        ]
    )
    ws.append(
        [
            "2026ME009",
            "Dev Mehta",
            "ME",
            "Software Engineering",
            15,
            "AB",
            16,
            30,
            None,
        ]
    )
    wb.save(path)


def test_excel_pipeline(tmp_path):
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    build_input(input_file)

    summary = process_workbook(str(input_file), str(output_file))

    wb = load_workbook(output_file)
    ws = wb.active

    assert summary.total_records == 4
    assert summary.valid_records == 1
    assert summary.invalid_records == 3
    assert summary.errors == [
        {
            "row": 3,
            "roll_number": "2026IT007",
            "message": "End examination marks out of range",
        },
        {"row": 4, "roll_number": "2026ECE008", "message": "Missing input"},
        {"row": 5, "roll_number": "2026ME009", "message": "Invalid data type"},
    ]
    assert ws.cell(2, 9).value == "C"
    assert ws.cell(2, 10).value is None
    assert ws.cell(3, 9).value is None
    assert ws.cell(3, 10).value == "End examination marks out of range"
    assert ws.cell(4, 9).value is None
    assert ws.cell(4, 10).value == "Missing input"
    assert ws.cell(5, 9).value is None
    assert ws.cell(5, 10).value == "Invalid data type"


def test_default_workbook_paths_are_anchored_to_lab_directory():
    assert (
        DEFAULT_INPUT_FILE
        == LAB_DIR / "data" / "input" / "Student_Records_Input.xlsx"
    )
    assert (
        DEFAULT_OUTPUT_FILE
        == LAB_DIR / "data" / "output" / "Student_Records_Updated.xlsx"
    )
